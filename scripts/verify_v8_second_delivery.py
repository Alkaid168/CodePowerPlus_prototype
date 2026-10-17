"""只读核验真实v8二测交付，验收收据不覆盖原预测封存。"""
from pathlib import Path
import hashlib
import importlib.util
import json
import math
import sys
from datetime import datetime, timezone
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.blind_evaluation import digest


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    base = ROOT / 'data/evaluation/v8-second-test'
    result = read(base / 'results.json')
    summary = read(base / 'summary.json')
    manifest = read(base / 'manifest.json')
    assert digest(result) == summary['results_sha256']
    for name, expected in manifest['fingerprint']['experiment_files'].items():
        assert sha(ROOT / name) == expected, name
    for name, expected in manifest['fingerprint']['original_fingerprint']['files'].items():
        assert sha(ROOT / name) == expected, name
    human = ROOT / 'output/v8-blind-form/v8首测-人工难度评测.xlsx'
    comparison = read(ROOT / 'data/evaluation/v8-first-test/comparison.json')
    assert sha(human) == comparison['human_file_sha256']
    output = ROOT / 'output/v8-second-test/v8二测-四次对比与代理复核.xlsx'
    audit = read(output.with_suffix('.audit.json'))
    assert audit['input_sha256'] == sha(base / 'results.json')
    assert audit['synthetic_fixture'] is False
    wb = load_workbook(output, data_only=True)
    formulas = load_workbook(output, data_only=False)
    assert wb.sheetnames == ['四次对比', '审核明细']
    main_sheet, detail = wb.worksheets
    keys = ['first', 'second_1', 'second_2', 'second_3']
    counts = []
    for index, item in enumerate(result['problems']):
        row = index + 12
        assert main_sheet.cell(row, 1).value == item['anonymous_id']
        assert main_sheet.cell(row, 2).value == item['title']
        assert main_sheet.cell(row, 3).hyperlink.target == item['url']
        expected = [item['human_rating']] + [item['runs'][key]['rating'] for key in keys]
        assert [main_sheet.cell(row, col).value for col in range(4, 9)] == expected
        count = sum(item['runs'][key]['joint_correct'] for key in keys)
        assert type(main_sheet.cell(row, 9).value) is int
        assert main_sheet.cell(row, 9).value == count
        assert formulas.worksheets[0].cell(row, 9).data_type == 'f'
        counts.append(count)
        for offset, key in enumerate(keys):
            record = item['runs'][key]
            r = index * 4 + offset + 5
            assert detail.cell(r, 1).value == item['anonymous_id']
            assert [detail.cell(r, c).value for c in [5, 6]] == [item['human_rating'], record['rating']]
            flags = [detail.cell(r, c).value for c in [7, 8, 9]]
            assert all(type(value) is bool for value in flags)
            assert flags == [record['tags_correct'], record['solution_correct'], record['joint_correct']]
            assert flags[2] == (flags[0] and flags[1])
            assert detail.cell(r, 12).value == record['evidence']
            assert detail.cell(r, 13).value == record['source_ref']
    for row, group in enumerate(['all20', 'regression9', 'new11'], 6):
        for col, key in enumerate(keys, 3):
            assert math.isclose(main_sheet.cell(row, col).value, summary['metrics'][key][group]['mae'])
    assert sum(counts) == 71 and len(counts) == 20
    assert main_sheet.cell(28, 4).value == 1500
    spec = importlib.util.spec_from_file_location('runner_delivery', ROOT / 'scripts/v8_second_test.py')
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    assert runner.runner().database_snapshot() == manifest['database_before']
    receipt = {'event': 'v8二测', 'state': 'complete', 'verified_at': datetime.now(timezone.utc).isoformat(),
        'new_analyses': 60, 'reviewed_analyses_including_first': 80, 'human_revision': 'r1',
        'joint_correct_by_run': [summary['metrics'][key]['all20']['joint_correct'] for key in keys],
        'mae_by_run': [summary['metrics'][key]['all20']['mae'] for key in keys],
        'workbook': str(output.relative_to(ROOT)).replace('\\', '/'), 'workbook_sha256': sha(output),
        'results_sha256': digest(result), 'sealed_predictions_sha256': summary['predictions_sha256'],
        'human_workbook_unchanged': True, 'production_database_unchanged': True,
        'prompts_and_frozen_implementation_unchanged': True, 'workbook_independent_readback': True,
        'previews_visually_inspected': True, 'native_excel_validation': False,
        'correctness_source': '代理逐份语义复核，非教师金标准或OJ AC率'}
    (base / 'outcome.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == '__main__':
    main()
