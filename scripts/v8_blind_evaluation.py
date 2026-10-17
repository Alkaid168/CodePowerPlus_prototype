"""v8首测冻结/真实预测/完整人工评分比较；不创建或连接可写正式数据库。"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import secrets
import sqlite3
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.blind_evaluation import compare, digest, validate_public, verify_commitment
from app.config import load_settings
from app.knowledge import Taxonomy
from app.services.analysis import analyze_with_model, PROMPT_VERSION

WORK = ROOT / 'data/evaluation/v8-first-test'
PRIVATE = Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'CodePowerPlus/blind-evaluations/v8-first-test'
REJECTED = {1013, 1014, 1017, 1033, 1034, 1035, 1036, 1037, 1039}


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + '.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    temp.replace(path)


def now():
    return datetime.now(timezone.utc).isoformat()


def database_snapshot():
    result = {}
    with sqlite3.connect((ROOT / 'data/codepowerplus.db').as_uri() + '?mode=ro', uri=True) as con:
        for table in ('problems', 'problem_analyses', 'analysis_reviews', 'submissions', 'schema_meta'):
            rows = con.execute(f'SELECT * FROM {table} ORDER BY rowid').fetchall()
            result[table] = {'count': len(rows), 'rows_sha256': digest(rows)}
    return result


def fingerprint(settings):
    paths = ['app/prompts/problem-analysis-v8-solve.txt', 'app/prompts/problem-analysis-v8-grade.txt',
             'app/services/analysis.py', 'app/services/analysis_v8.py', 'app/schemas.py',
             'app/services/difficulty.py', 'app/knowledge.py', 'data/knowledge_taxonomy.json',
             'app/blind_evaluation.py', 'scripts/v8_blind_evaluation.py', 'docs/v8-first-test-protocol.md']
    hashes = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in paths}
    return {'files': hashes, 'prompt_version': PROMPT_VERSION,
            'request_model': settings.model, 'endpoint': settings.base_url, 'temperature': 0,
            'request_timeout': settings.request_timeout}


def freeze(settings):
    path = WORK / 'manifest.json'
    if path.exists():
        frozen = read(path)
        if frozen['fingerprint'] != fingerprint(settings):
            raise ValueError('实现/提示词/模型/协议已变动，不能悄悄复用同一次实验')
        return frozen
    old = read(ROOT / 'data/evaluation/reassess-2026-10-17-v7/manifest.json')['items']
    items = []
    for row in old:
        if row['id'] in REJECTED:
            items.append({'group': 'regression9', 'origin_problem_id': row['id'],
                          'source_id': str(row['id']), 'source_url': row['source_url'],
                          'description': row['description'], 'reference_solution': row['reference_solution'],
                          'reference_code': row['reference_code'], 'input_origin': 'v7_frozen_exact'})
    new = read(WORK / 'new11.json')['items']
    for row in new:
        items.append({key: row[key] for key in ('source_id', 'source_url', 'description', 'reference_solution', 'reference_code')}
                     | {'group': 'new11', 'input_origin': 'official_verified_independent_restatement'})
    if len(items) != 20 or len(new) != 11 or {x.get('origin_problem_id') for x in items if x['group'] == 'regression9'} != REJECTED:
        raise ValueError('样本必须精确覆盖9道驳回题与11道新增题')
    secrets.SystemRandom().shuffle(items)
    for i, row in enumerate(items, 1):
        row['public_id'] = f'B{i:02}'
        row['input_sha256'] = digest({k: row[k] for k in ('description', 'reference_solution', 'reference_code')})
    public = {'items': [{k: row[k] for k in ('public_id', 'source_url')} for row in items]}
    validate_public(public)
    frozen = {'event': 'v8首测', 'frozen_at': now(), 'fingerprint': fingerprint(settings),
              'database_before': database_snapshot(), 'items': items}
    write(path, frozen)
    write(WORK / 'public_manifest.json', public)
    return frozen


def predict(settings):
    frozen = freeze(settings)
    manifest_hash = digest(frozen)
    checkpoint_path = PRIVATE / 'checkpoint.json'
    sealed_path = PRIVATE / 'predictions.json'
    if (WORK / 'commitment.json').exists():
        seal = read(WORK / 'commitment.json')
        verify_commitment(read(sealed_path), seal)
        if seal['manifest_sha256'] != manifest_hash:
            raise ValueError('封存输入已变动')
        print('20/20 已封存，等待人工评分；没有重新调用模型。')
        return
    if not settings.model_configured:
        raise ValueError('未配置真实模型密钥，不能开始预测')
    checkpoint = read(checkpoint_path) if checkpoint_path.exists() else {
        'manifest_sha256': manifest_hash, 'started_at': now(), 'results': {}, 'failures': []}
    if checkpoint['manifest_sha256'] != manifest_hash:
        raise ValueError('断点与冻结清单不匹配')
    taxonomy = Taxonomy(settings.taxonomy_path, settings.schema_path)
    def run(row):
        result = analyze_with_model(row['description'], taxonomy, settings,
                    reference_solution=row['reference_solution'], reference_code=row['reference_code'])
        return {'public_id': row['public_id'], 'group': row['group'],
                'input_sha256': row['input_sha256'], 'analysis': result.model_dump()}
    pending = [x for x in frozen['items'] if x['public_id'] not in checkpoint['results']]
    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = {executor.submit(run, row): row for row in pending}
        for future in as_completed(futures):
            row = futures[future]
            try:
                checkpoint['results'][row['public_id']] = future.result()
                print(f"{row['public_id']} 完成；累计 {len(checkpoint['results'])}/20", flush=True)
            except Exception as exc:
                checkpoint['failures'].append({'public_id': row['public_id'], 'at': now(), 'error_type': type(exc).__name__})
                print(f"{row['public_id']} 暂未完成（{type(exc).__name__}），保留失败记录", flush=True)
            write(checkpoint_path, checkpoint)
    if len(checkpoint['results']) != 20:
        raise ValueError('预测未齐20题：尚未封存或发放人工文件；可原样断点重试失败项')
    predictions = {'event': 'v8首测', 'manifest_sha256': manifest_hash, 'completed_at': now(),
        'started_at': checkpoint['started_at'], 'failures': checkpoint['failures'],
        'items': [checkpoint['results'][x['public_id']] for x in frozen['items']]}
    if database_snapshot() != frozen['database_before']:
        raise ValueError('正式数据库在实验期间发生变化，请核查，暂不封存')
    write(sealed_path, predictions)
    write(WORK / 'commitment.json', {'event': 'v8首测', 'state': 'awaiting_human', 'completed': 20,
        'sealed_at': now(), 'manifest_sha256': manifest_hash, 'predictions_sha256': digest(predictions),
        'public_manifest_sha256': digest(read(WORK / 'public_manifest.json')),
        'database_unchanged': True, 'confidence_is_calibrated': False,
        'metric_protocol': 'docs/v8-first-test-protocol.md'})
    print('20/20 预测封存完成，公开承诺哈希已写入；等待人工评分。', flush=True)


def human_rows(path):
    if path.suffix.lower() == '.json':
        return read(path)
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False, read_only=True)
    if workbook.sheetnames != ['人工难度']:
        raise ValueError('人工文件工作表不匹配')
    sheet = workbook['人工难度']
    if sheet.max_row != 24 or sheet.max_column != 3:
        raise ValueError('人工文件行列范围已改变')
    return [{'public_id': sheet.cell(i, 1).value, 'source_url': sheet.cell(i, 2).value,
             'difficulty': sheet.cell(i, 3).value} for i in range(5, 25)]


def reveal(path):
    frozen = read(WORK / 'manifest.json')
    public = read(WORK / 'public_manifest.json')
    seal = read(WORK / 'commitment.json')
    predictions = read(PRIVATE / 'predictions.json')
    verify_commitment(predictions, seal)
    if digest(frozen) != seal['manifest_sha256'] or digest(public) != seal['public_manifest_sha256']:
        raise ValueError('冻结清单哈希不匹配')
    # 完整且有效的人评到达后，才将预测写进可见报告。
    report = compare(public, predictions, human_rows(path))
    report['human_file_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    report['prediction_commitment'] = seal['predictions_sha256']
    destination = WORK / 'comparison.json'
    if destination.exists() and read(destination) != report:
        raise ValueError('已有不同人工比较结果，不能覆盖；需新建复评事件')
    write(destination, report)
    lines = ['# v8首测难度比较', '', '差值 = 模型 − 人工；负值表示模型低估。', '',
             '| 组别 | 题数 | MAE | 中位绝对差 | 有向均差 | ≤100 | ≤200 | ≤300 | 最大绝对差 |',
             '| --- | --- | --- | --- | --- | --- | --- | --- | --- |']
    for group, m in report['metrics'].items():
        lines.append(f"| {group} | {m['n']} | {m['mae']:.1f} | {m['median_absolute_difference']:.1f} | {m['bias']:.1f} | {m['within_100']:.1%} | {m['within_200']:.1%} | {m['within_300']:.1%} | {m['max_absolute_difference']} |")
    lines += ['', '| ID | 原题 | 模型 | 人工 | 差值 |', '| --- | --- | --- | --- | --- |']
    for row in report['items']:
        lines.append(f"| {row['public_id']} | [原题]({row['source_url']}) | {row['model_difficulty']} | {row['human_difficulty']} | {row['difference']:+} |")
    lines += ['', report['limitations']]
    (WORK / 'comparison.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print('完整人工评分校验通过，分差与分组统计已保存。')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['freeze', 'predict', 'compare'])
    parser.add_argument('--human', type=Path)
    args = parser.parse_args()
    if args.action == 'compare':
        if args.human is None:
            parser.error('比较需要--human人工文件')
        reveal(args.human.resolve())
    elif args.action == 'freeze':
        freeze(load_settings())
        print('20题输入冻结完成，没有调用模型。')
    else:
        predict(load_settings())
